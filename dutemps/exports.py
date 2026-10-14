"""CSV and XLSX exports. The detailed export has a fixed set of 18 columns, so
anything built on it keeps working."""
import csv
import io

from . import core, db, reports
from .timeutil import hms, local

DETAILED_COLUMNS = [
    "Project", "Client", "Description", "Task", "User", "Group", "Email", "Tags",
    "Billable", "Start Date", "Start Time", "End Date", "End Time", "Duration (h)",
    "Duration (decimal)", "Billable Rate (USD)", "Billable Amount (USD)",
    "Date of creation",
]


def detailed_rows(con, f):
    s = db.settings(con)
    tz = s["timezone"]
    out = []
    for e in core.list_entries(con, {**f, "finished": True}, order="ASC"):
        a, b = local(e["start_utc"], tz), local(e["end_utc"], tz)
        out.append([
            e["project"] or "", e["client"] or "", e["description"], e["task"] or "",
            s["user_name"], "", s["user_email"],
            ", ".join(t["name"] for t in e["tags"]),
            "Yes" if e["billable"] else "No",
            a.strftime("%Y-%m-%d"), a.strftime("%H:%M:%S"),
            b.strftime("%Y-%m-%d"), b.strftime("%H:%M:%S"),
            hms(e["seconds"]), round(e["seconds"] / 3600, 4),
            round(e["rate"], 2) if e["billable"] and e["rate"] is not None else 0,
            e["amount"],
            local(e["created_utc"], tz).strftime("%Y-%m-%d"),
        ])
    return DETAILED_COLUMNS, out


def summary_rows(con, f, group1, group2):
    r = reports.summary(con, f, group1, group2)
    head = [group1.title()] + ([group2.title()] if group2 else []) + [
        "Duration (h)", "Duration (decimal)", "Billable (decimal)", "Amount (USD)"]
    rows = []
    for g in r["groups"]:
        kids = g["children"] if group2 else [None]
        for c in kids:
            x = c or g
            rows.append([g["name"]] + ([c["name"]] if group2 else []) + [
                hms(x["seconds"]), x["hours"], x["billable_hours"], x["amount"]])
    t = r["total"]
    rows.append(["Total"] + ([""] if group2 else []) +
                [hms(t["seconds"]), t["hours"], t["billable_hours"], t["amount"]])
    return head, rows


def weekly_rows(con, f, monday):
    w = reports.weekly(con, f, monday)
    head = ["Project", "Task"] + w["days"] + ["Total (decimal)"]
    rows = [[r["project"], r["task"] or ""] + [round(x / 3600, 4) for x in r["days"]]
            + [round(r["total"] / 3600, 4)] for r in w["rows"]]
    rows.append(["Total", ""] + [round(x / 3600, 4) for x in w["totals"]]
                + [round(w["total"] / 3600, 4)])
    return head, rows


def invoice_rows(con, f):
    r = reports.invoice_lines(con, f)
    head = ["Date", "Project", "Client", "Task", "Description", "Raw (decimal)",
            "Billed (decimal)", "Rate (USD)", "Amount (USD)", "Entries"]
    rows = [[l["date"], l["project"], l["client"], l["task"], l["description"],
             l["raw_hours"], l["hours"], l["rate"], l["amount"], l["entries"]]
            for l in r["lines"]]
    rows.append(["Total", "", "", "", "", r["raw_hours"], r["hours"], "", r["amount"], ""])
    return head, rows


def table(con, kind, f, group1="project", group2=None, monday=None):
    if kind == "detailed":
        return detailed_rows(con, f)
    if kind == "summary":
        return summary_rows(con, f, group1, group2)
    if kind == "weekly":
        return weekly_rows(con, f, monday)
    if kind == "invoice":
        return invoice_rows(con, f)
    raise core.Invalid(f"unknown report {kind}")


def to_csv(head, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(head)
    w.writerows(rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")     # BOM: Excel reads UTF-8


def to_xlsx(head, rows, title):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise core.Invalid("Excel export needs openpyxl (pip install openpyxl); "
                           "CSV export works without it") from None
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(head, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
